"""Prints a JSON description of an .xlsx or .pdf file for task graders.

    python3 -I file_inspect.py xlsx book.xlsx   -> {"sheets": {"Name": [[{"v":..., "f":..., "fmt":...}, ...], ...]}}
    python3 -I file_inspect.py pdf file.pdf     -> {"pages": 3, "text": ["page 1 text", ...]}

Exit code 2 means the file could not be opened as that format.
"""
import json
import sys


def xlsx(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, data_only=False)
    sheets = {}
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows():
            cells = []
            for cell in row:
                value = cell.value
                formula = value if isinstance(value, str) and value.startswith("=") else None
                if formula is None and hasattr(value, "text"):  # ArrayFormula
                    formula = value.text
                cells.append({"v": None if formula else value, "f": formula, "fmt": cell.number_format})
            rows.append(cells)
        sheets[sheet.title] = rows
    return {"sheets": sheets}


def pdf(path):
    from pypdf import PdfReader
    reader = PdfReader(path)
    return {"pages": len(reader.pages), "text": [page.extract_text() or "" for page in reader.pages]}


if __name__ == "__main__":
    kind, path = sys.argv[1], sys.argv[2]
    try:
        result = xlsx(path) if kind == "xlsx" else pdf(path)
    except Exception as error:  # noqa: BLE001 - any failure means the file is unusable
        print(json.dumps({"error": f"{type(error).__name__}: {error}"}))
        sys.exit(2)
    print(json.dumps(result, default=str))
